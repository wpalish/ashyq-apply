"""Accepted evidence survives XLSX save/reload without becoming a formula.

Also: each exported row names where its facts come from (EX-1).
"""

from __future__ import annotations

import csv
import io
from datetime import UTC, date, datetime

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.catalogue.importer import import_catalogue
from app.catalogue.retrieval import retrieve
from app.db import create_all_for_tests
from app.domain.enums import ClaimType
from app.export import tabular
from app.export.tabular import COLUMNS, KB_EXTRACTION_METHOD, to_csv, to_xlsx
from app.models import ProgramResultRow, ResearchRun
from app.pipeline.runner import ResearchRunner
from app.pipeline.state import RunState
from app.schemas.claim import MAX_EXCERPT_CHARS, ClaimOut
from app.schemas.result import ProgramResult
from tests.conftest import make_claim, profile_row
from tests.test_knowledge_base import _kb_settings

_CONDITION = (
    "The scholarship covers up to the normal programme duration on condition that the "
    "scholarship holder maintain a record of good academic performance and exemplary conduct."
)
_POLICY = {
    "conditions": [_CONDITION],
    "tuition": {"fraction": 1.0, "basis": "subsidised_after_tuition_grant"},
    "living": {"currency": "SGD", "fixed_amount": 6500, "period": "academic_year"},
    "housing": {
        "currency": "SGD",
        "maximum": 2000,
        "period": "academic_year",
        "requires": "scholarship holders who reside in NTU hostels only",
    },
    "travel": {
        "currency": "SGD",
        "maximum": 8000,
        "requires": "an overseas programme",
        "terms": "Travel Grant Form",
        "cohorts_from": "AY2025",
    },
}
_ACTUAL_QUOTE = (
    _CONDITION + "\nFull coverage of subsidised tuition fees (after Tuition Grant)."
    "\nLiving allowance of S$6,500 per academic year."
    "\nAccommodation allowance of up to S$2,000 per academic year."
    "\n(Applicable to scholarship holders who reside in NTU hostels only.)"
    "\nTravel grant of up to S$8,000 for an overseas programme subject to terms and conditions "
    "in the Travel Grant Form (for new cohorts from AY2025)."
)
_CAP_QUOTE = "Published conditions: " + "x" * 549 + " Last condition must survive."
_URL = (
    "https://www.ntu.edu.sg/admissions/undergraduate/scholarships/"
    "scholarship-opportunities/detail/nanyang-scholarship"
)
_HEADERS = [
    "University",
    "Program",
    "Claim type",
    "Value",
    "Status",
    "Specificity",
    "Accessed",
    "Source URL",
    "Excerpt",
]


def _result(value: object, excerpt: str) -> ProgramResult:
    claim = make_claim(
        ClaimType.SCHOLARSHIP_COVERAGE.value,
        value,
        url=_URL,
        specificity="scholarship_administrator",
        subject_key="Nanyang Global Scholarship",
        accessed_at=datetime(2026, 9, 29, 8, 30, 34, 955612, tzinfo=UTC),
        academic_year="2026/27",
    )
    payload = claim.model_dump()
    payload["original_text_excerpt"] = excerpt
    return ProgramResult(
        id="result-1",
        run_id="run-1",
        university="Nanyang Technological University",
        university_id="ntu",
        country="Singapore",
        city="Singapore",
        program="Bachelor Of Computing In Computer Science",
        degree="bachelor",
        intake="Fall 2027",
        claims=[ClaimOut(id="claim-1", **payload)],
    )


def test_complete_nested_policy_survives_workbook_round_trip():
    result = _result(_POLICY, _ACTUAL_QUOTE)
    workbook = load_workbook(io.BytesIO(to_xlsx([result])))

    assert workbook.sheetnames == ["Shortlist", "Evidence", "Open questions"]
    evidence = workbook["Evidence"]
    assert [cell.value for cell in evidence[1]] == _HEADERS
    expected = str(result.claims[0].normalized_value)
    assert len(expected) > 200
    assert evidence["D2"].value == expected
    assert "Travel Grant Form" in evidence["D2"].value
    assert "AY2025" in evidence["D2"].value
    assert "scholarship holders who reside in NTU hostels only" in evidence["D2"].value
    assert evidence["H2"].value == _URL


@pytest.mark.parametrize("excerpt", [_ACTUAL_QUOTE, _CAP_QUOTE], ids=["actual-552", "cap-600"])
def test_complete_accepted_excerpt_survives_workbook_round_trip(excerpt):
    assert len(_ACTUAL_QUOTE) == 552
    assert len(_CAP_QUOTE) == MAX_EXCERPT_CHARS
    result = _result(_POLICY, excerpt)
    workbook = load_workbook(io.BytesIO(to_xlsx([result])))

    assert result.claims[0].original_text_excerpt == excerpt
    assert workbook["Evidence"]["I2"].value == excerpt


@pytest.mark.parametrize("lead", ["=", "+", "-", "@"])
def test_long_untrusted_evidence_stays_complete_and_inert(lead):
    value = lead + 'HYPERLINK("https://attacker.invalid/", "' + "x" * 480 + '")'
    result = _result(value, value)
    workbook = load_workbook(io.BytesIO(to_xlsx([result])), data_only=False)
    evidence = workbook["Evidence"]

    for cell in (evidence["D2"], evidence["I2"]):
        assert cell.value == "'" + value
        assert cell.data_type == "s"


# ------------------------------------------------------------- data origin (EX-1)

_KB_ORIGIN = "verified knowledge base (official sources, checked 2026-10-19)"
_ORIGIN_COLUMN = [key for key, _ in COLUMNS].index("data_origin")


def _old_data_origin(result: ProgramResult) -> str:
    """The column as it was before EX-1: two values, the oracle for live and fixture rows."""
    return (
        "DEMO FIXTURE (synthetic)"
        if any(u.startswith("fixture://") for u in result.source_urls)
        else "live official sources"
    )


def _csv_rows(text: str) -> list[list[str]]:
    """The header, then one list per row (the disclaimer line dropped)."""
    return list(csv.reader(io.StringIO(text)))[1:]


def _xlsx_origins(data: bytes) -> list[str]:
    sheet = load_workbook(io.BytesIO(data))["Shortlist"]
    return [row[_ORIGIN_COLUMN].value for row in sheet.iter_rows(min_row=3)]


@pytest.fixture
def kb_session(monkeypatch):
    monkeypatch.setattr("app.knowledge.claims.kb_today", lambda: date(2026, 10, 19))
    engine = create_engine("sqlite://")
    create_all_for_tests(engine)
    with Session(engine) as session:
        import_catalogue(session)
        session.commit()
        yield session
    engine.dispose()


async def _kb_results(session: Session, profile, settings, *queries: str) -> list[ProgramResult]:
    """A kb-mode run over the named universities, its rows as the export reads them."""
    ids = [retrieve(session, query=q)[0].id for q in queries]
    run = ResearchRun(
        profile_id=profile_row(session, profile).id,
        stage="queued",
        demo_mode=False,
        candidate_limit=len(ids),
        verify_limit=len(ids),
        university_ids=ids,
        stage_state=RunState.load(None).dump(),
    )
    session.add(run)
    session.flush()
    await ResearchRunner(session, run, profile, settings).run_to_decision()
    rows = session.scalars(select(ProgramResultRow).where(ProgramResultRow.run_id == run.id))
    return [ProgramResult.model_validate(r.payload) for r in rows]


@pytest.mark.asyncio
async def test_a_kb_row_names_the_knowledge_base_and_its_date_in_csv_and_xlsx(
    kb_session, profile, settings, tmp_path
):
    (result,) = await _kb_results(
        kb_session, profile, _kb_settings(settings, tmp_path), "Massachusetts Institute"
    )
    # The field the decision reads: every KB claim carries it, and the date is
    # the row's own last_verified (the KB's verified_on).
    assert result.claims
    assert {c.extraction_method for c in result.claims} == {KB_EXTRACTION_METHOD}
    assert result.last_verified is not None
    assert result.last_verified.date() == date(2026, 10, 19)

    header, row = _csv_rows(to_csv([result]))
    assert header == [label for _, label in COLUMNS]
    assert len(row) == len(COLUMNS) == 38
    assert row[_ORIGIN_COLUMN] == _KB_ORIGIN
    assert _xlsx_origins(to_xlsx([result])) == [_KB_ORIGIN]


def _plain_row(url: str, method: str = "html_rule") -> ProgramResult:
    claim = make_claim(
        ClaimType.IELTS_MIN_OVERALL.value,
        6.5,
        url=url,
        accessed_at=datetime(2026, 9, 29, 8, 30, tzinfo=UTC),
    ).model_copy(update={"extraction_method": method})
    return ProgramResult(
        id=f"result-{method}",
        run_id="run-1",
        university="Example University",
        university_id="example",
        country="Netherlands",
        city="Groningen",
        program="BSc Computing Science",
        degree="bachelor",
        intake="Fall 2027",
        claims=[ClaimOut(id="claim-1", **claim.model_dump())],
        source_urls=[url],
        last_verified=datetime(2026, 9, 29, 8, 30, tzinfo=UTC),
    )


def test_live_and_fixture_rows_export_exactly_as_before(monkeypatch):
    live = _plain_row("https://www.example.edu/programme")
    fixture = _plain_row("fixture://example/programme", method="fixture")
    results = [live, fixture]

    csv_now, xlsx_origins_now = to_csv(results), _xlsx_origins(to_xlsx(results))
    monkeypatch.setattr(tabular, "data_origin", _old_data_origin)
    csv_before, xlsx_origins_before = to_csv(results), _xlsx_origins(to_xlsx(results))

    assert csv_now == csv_before  # header and every value, byte for byte
    assert xlsx_origins_now == xlsx_origins_before
    assert [r[_ORIGIN_COLUMN] for r in _csv_rows(csv_now)[1:]] == [
        "live official sources",
        "DEMO FIXTURE (synthetic)",
    ]
    assert _csv_rows(csv_now)[0] == [label for _, label in COLUMNS]


@pytest.mark.asyncio
async def test_a_mixed_kb_run_labels_each_row_by_its_own_claims(
    kb_session, profile, settings, tmp_path
):
    results = await _kb_results(
        kb_session, profile, _kb_settings(settings, tmp_path), "Massachusetts Institute", "Stanford"
    )
    by_name = {r.university: r for r in results}
    (mit,) = [r for name, r in by_name.items() if "Massachusetts" in name]
    (stanford,) = [r for name, r in by_name.items() if "Stanford" in name]
    assert stanford.claims == []
    assert any(q.question.startswith("No verified data yet") for q in stanford.unresolved)

    ordered = [mit, stanford]
    origins = [row[_ORIGIN_COLUMN] for row in _csv_rows(to_csv(ordered))[1:]]
    assert origins == [_KB_ORIGIN, "live official sources"]
    assert _xlsx_origins(to_xlsx(ordered)) == origins
